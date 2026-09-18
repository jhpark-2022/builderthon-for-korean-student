# 명단 문서 공통 부분 (2026-09-18, Supabase 등록 브리프 2.6).
# build-applicant-roster.py(8월)와 build-crossing-roster.py(12월)가 같이 씁니다.
# 8월 스크립트에서 옮겼고, 동작은 그대로입니다(리팩토링 전후 같은 docx).

import json
import subprocess
import sys
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path

from docx import Document

REPO = Path(__file__).resolve().parent.parent
KST = timezone(timedelta(hours=9))
EM = " "  # 전각 공백(U+2003 em space). 문서 전체가 이걸로 항목을 띄운다.
NO_ENTRY = "미기재"


def load_env():
    env = {}
    for line in (REPO / ".env.local").read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        env[k.strip()] = v.strip()
    url = env.get("NEXT_PUBLIC_SUPABASE_URL")
    key = env.get("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        sys.exit(".env.local에 NEXT_PUBLIC_SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY가 필요합니다.")
    return url.rstrip("/"), key


def fetch(url, key, path):
    # 시스템 파이썬에 CA 번들이 없어 urllib이 Supabase TLS를 검증하지 못하는 경우가 있다.
    # curl은 macOS 키체인을 쓰므로 그쪽으로 붙는다.
    out = subprocess.run(
        [
            "curl", "-sS", "--fail", f"{url}/rest/v1/{path}",
            "-H", f"apikey: {key}",
            "-H", f"Authorization: Bearer {key}",
        ],
        capture_output=True,
        text=True,
    )
    if out.returncode != 0:
        sys.exit(f"Supabase 요청 실패 ({path}): {out.stderr.strip()}")
    return json.loads(out.stdout)


def kst(iso, fmt="%Y-%m-%d %H:%M"):
    # Supabase는 마이크로초 자리수가 들쭉날쭉해서 fromisoformat에 그대로 넣기 어렵다.
    # 소수점 이하를 정확히 6자리로 맞춰 준다.
    s = iso.replace("Z", "+00:00")
    if "." in s:
        head, rest = s.split(".", 1)
        digits = ""
        while rest and rest[0].isdigit():
            digits, rest = digits + rest[0], rest[1:]
        s = f"{head}.{digits[:6].ljust(6, '0')}{rest}"
    return datetime.fromisoformat(s).astimezone(KST).strftime(fmt)


def ranked(pairs):
    """(값, 최초등장순번) 목록을 건수 내림차순 → 최초 등장 순으로 정렬한다."""
    counts = Counter(v for v, _ in pairs)
    first = {}
    for v, seq in pairs:
        first.setdefault(v, seq)
    return sorted(counts.items(), key=lambda kv: (-kv[1], first[kv[0]]))


def clear_body(doc):
    body = doc.element.body
    for child in list(body.iterchildren()):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, headers):
        cell.paragraphs[0].add_run(text).bold = True
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = "" if text is None else str(text)
    doc.add_paragraph("")
    return t


def open_doc(out: Path):
    """기존 파일을 서식 템플릿으로 열어 본문만 비운다. 없으면 새 문서."""
    if out.exists():
        doc = Document(str(out))
        clear_body(doc)
        return doc
    out.parent.mkdir(parents=True, exist_ok=True)
    return Document()
