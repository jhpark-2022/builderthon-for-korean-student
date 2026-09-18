#!/usr/bin/env python3
# 빌더톤_신청자_명단.docx 생성기.
#
# Supabase의 registrations / registration_members를 읽어 운영용 명단 문서를 다시 만든다.
# 기존 파일을 서식 템플릿으로 열어 본문만 비우고 채우므로, 글꼴·표 스타일·용지 설정이
# 회차마다 흔들리지 않는다.
#
#   python3 scripts/build-applicant-roster.py
#
# 자격증명은 website/.env.local의 NEXT_PUBLIC_SUPABASE_URL + SUPABASE_SERVICE_ROLE_KEY.
# service_role 키라 RLS를 우회한다. 절대 브라우저로 넘기지 말 것.

import os
import sys
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from docx import Document

# 2026-09-18: 공통 부분(env, REST fetch, KST, 정렬, 표)은 scripts/roster_lib.py로 꺼냈습니다.
# 12월 스크립트(build-crossing-roster.py)와 같이 씁니다. 이 스크립트의 결과 docx는 그대로입니다.
sys.path.insert(0, str(Path(__file__).resolve().parent))
from roster_lib import REPO, KST, EM, NO_ENTRY, load_env, fetch, kst, ranked, clear_body, add_table  # noqa: E402

# ROSTER_OUT을 주면 그 경로에 씁니다(회귀 확인용). 기본은 Dropbox의 운영 문서.
OUT = Path(os.environ["ROSTER_OUT"]) if os.environ.get("ROSTER_OUT") else REPO.parent / "Execution" / "Tracking" / "빌더톤_신청자_명단.docx"

TRACK_LABELS = {
    "unsure": "아직 모르겠음",
    "ops_automation": "업무 자동화",
    "engineering": "개발 엔지니어링",
    "data_analytics": "데이터 분석",
    "customer_marketing": "고객 마케팅",
    "finance": "금융(구 문항)",
    "sales": "영업(구 문항)",
    "marketing": "마케팅(구 문항)",
}


def track_label(track):
    if not track:
        return NO_ENTRY
    return TRACK_LABELS.get(track, track)


def main():
    url, key = load_env()
    regs = fetch(url, key, "registrations?select=*&order=created_at.asc")
    members = fetch(url, key, "registration_members?select=*")

    by_reg = {}
    for m in members:
        by_reg.setdefault(m["registration_id"], []).append(m)

    # 신청번호(1..N)와 사람 단위 행을 한 번에 만든다.
    people = []  # 문서 3장의 한 행 = 한 사람
    for no, r in enumerate(regs, 1):
        r["no"] = no
        r["members"] = sorted(by_reg.get(r["id"], []), key=lambda m: m["ordinal"])
        for m in r["members"]:
            m["reg"] = r
            m["role"] = "대표" if m["ordinal"] == 1 else f"팀원{m['ordinal'] - 1}"
            people.append(m)

    # ── 중복 신청 찾기: 서로 다른 신청서에 같은 이메일 또는 같은 연락처가 있는 경우.
    # 김승학 건처럼 이메일에 오타(gmai.com)가 있어도 연락처로 잡히도록 두 축을 모두 본다.
    groups = OrderedDict()  # 대표키 -> {"members": [...], "by": "이메일"|"연락처", "value": str}
    seen_email, seen_contact = {}, {}
    for p in people:
        email = (p.get("email") or "").strip().lower()
        contact = (p.get("contact") or "").strip().lower()
        hit = None
        if email and email in seen_email:
            hit, by, value = seen_email[email], "이메일", email
        elif contact and contact in seen_contact:
            hit, by, value = seen_contact[contact], "연락처", contact
        if hit is not None and hit["reg"]["no"] != p["reg"]["no"]:
            g = groups.setdefault(id(hit), {"members": [hit], "by": by, "value": value})
            g["members"].append(p)
        seen_email.setdefault(email, p)
        seen_contact.setdefault(contact, p)

    # ── 중복 해소: 그룹마다 가장 나중 신청서를 남기고 나머지는 명단에서 뺀다.
    # 나중 것이 본인의 최신 의사(매칭 희망 변경, 유형 재응시 등)를 담고 있다고 본다.
    # 뺀 기록은 2장에 그대로 남겨 무엇을 어떤 근거로 제외했는지 되짚을 수 있게 한다.
    def snap(m, state):
        r = m["reg"]
        return [
            state,
            kst(r["created_at"]),
            "팀" if r.get("join_type") == "team" else "개인",
            r.get("team_name") or "-",
            m["role"],
            m["name"],
            m["email"],
            m["contact"],
            m.get("university") or NO_ENTRY,
            "예" if r.get("wants_matching") else "아니오",
            r.get("quiz_type") or "-",
        ]

    dropped = []
    dropped_ids = set()
    for g in groups.values():
        keep = max(g["members"], key=lambda m: (m["reg"]["created_at"], m["reg"]["no"]))
        for m in g["members"]:
            if m is keep:
                continue
            dropped_ids.add(id(m))
            dropped.append(
                {
                    "name": m["name"],
                    "by": g["by"],
                    "value": g["value"],
                    "rows": [snap(m, "제외"), snap(keep, "유지")],
                }
            )

    # 사람을 빼고 나서 신청서가 통째로 비면 그 신청서도 없앤다(1인 개인 신청의 중복 건).
    for r in regs:
        r["members"] = [m for m in r["members"] if id(m) not in dropped_ids]
    regs = [r for r in regs if r["members"]]
    for no, r in enumerate(regs, 1):
        r["no"] = no
    people = [m for r in regs for m in r["members"]]

    # ── 통계
    solo = sum(1 for r in regs if r.get("join_type") != "team")
    team = len(regs) - solo
    matching = sum(1 for r in regs if r.get("wants_matching"))
    tracks = ranked([(track_label(r.get("track")), r["no"]) for r in regs])
    unis = ranked([((p.get("university") or "").strip() or NO_ENTRY, i) for i, p in enumerate(people)])
    refs = ranked([(r.get("ref") or "직접 방문", r["no"]) for r in regs])
    quizzes = ranked([(r["quiz_type"], r["no"]) for r in regs if r.get("quiz_type")])

    doc = Document(str(OUT))
    clear_body(doc)

    doc.add_paragraph("한인 빌더톤 신청자 명단", style="Title")
    now = datetime.now(KST).strftime("%Y년 %m월 %d일 %H:%M")
    head = f"추출 시각: {now} (KST){EM}신청 건수 {len(regs)}건{EM}명단 인원 {len(people)}명"
    if dropped:
        head += f"{EM}중복 제외 {len(dropped)}명"
    doc.add_paragraph(head)

    # ── 1. 요약 통계
    doc.add_paragraph("1. 요약 통계", style="Heading 1")
    add_table(
        doc,
        ["구분", "값"],
        [
            ["총 신청 건수", f"{len(regs)}건 (중복 제외 후)"],
            ["명단 인원", f"{len(people)}명 (전원 고유 인원)"],
            ["중복 제외", f"{len(dropped)}명 (아래 2장 참고)" if dropped else "없음"],
            ["참가 유형", f"개인 {solo}건, 팀 {team}건"],
            ["팀 매칭 희망", f"{matching}건"],
            ["트랙", ", ".join(f"{k} {v}건" for k, v in tracks)],
            ["소속 대학", ", ".join(f"{k} {v}명" for k, v in unis)],
            ["유입 경로", ", ".join(f"{k} {v}건" for k, v in refs)],
            ["최초 신청", kst(regs[0]["created_at"])],
            ["최근 신청", kst(regs[-1]["created_at"])],
        ],
    )

    doc.add_paragraph("퀴즈 유형 분포", style="Heading 2")
    add_table(doc, ["유형", "인원"], [[k, f"{v}건"] for k, v in quizzes])

    # ── 2. 중복 신청
    doc.add_paragraph("2. 중복 제외 내역", style="Heading 1")
    if dropped:
        doc.add_paragraph(
            f"같은 사람이 두 번 신청한 {len(dropped)}건을 명단에서 제외했습니다. "
            "나중 신청서를 본인의 최신 의사로 보고 그쪽을 남겼습니다. "
            "1장의 집계와 3장 이후의 모든 표는 제외를 마친 기준입니다."
        )
        for i, d in enumerate(dropped, 1):
            doc.add_paragraph(
                f"제외 {i}{EM}{d['name']} ({d['by']} 일치: {d['value']})", style="Heading 2"
            )
            add_table(
                doc,
                ["상태", "신청일시", "유형", "팀명", "역할", "이름", "이메일",
                 "연락처(카톡ID)", "대학", "매칭 희망", "퀴즈 유형"],
                d["rows"],
            )
    else:
        doc.add_paragraph("같은 사람이 두 번 신청한 건은 없습니다.")
        doc.add_paragraph("")

    # ── 3. 참가자 전체 명단
    doc.add_paragraph("3. 참가자 전체 명단 (1인 1행)", style="Heading 1")
    doc.add_paragraph("중복은 이미 제외했습니다. 비고는 현장에서 손으로 적기 위해 비워 둔 칸입니다.")
    add_table(
        doc,
        ["No", "신청번호", "유형", "팀명", "역할", "이름", "이메일", "연락처(카톡ID)", "대학", "트랙", "신청일시", "비고"],
        [
            [
                i,
                p["reg"]["no"],
                "팀" if p["reg"].get("join_type") == "team" else "개인",
                p["reg"].get("team_name") or "-",
                p["role"],
                p["name"],
                p["email"],
                p["contact"],
                p.get("university") or "",
                track_label(p["reg"].get("track")),
                kst(p["reg"]["created_at"]),
                "",
            ]
            for i, p in enumerate(people, 1)
        ],
    )

    # ── 4. 신청 건별 상세
    doc.add_paragraph("4. 신청 건별 상세", style="Heading 1")
    for r in regs:
        is_team = r.get("join_type") == "team"
        title = r.get("team_name") or (r["members"][0]["name"] if r["members"] else "-")
        head = f"[{r['no']}] {title}{EM}{'팀' if is_team else '개인'}{EM}{len(r['members'])}명"
        doc.add_paragraph(head, style="Heading 2")
        doc.add_paragraph(
            f"신청일시 {kst(r['created_at'])} (KST){EM}|{EM}트랙 {track_label(r.get('track'))}"
            f"{EM}|{EM}매칭 희망 {'예' if r.get('wants_matching') else '아니오'}"
            f"{EM}|{EM}퀴즈 유형 {r.get('quiz_type') or '-'}"
            f"{EM}|{EM}유입 {r.get('ref') or '직접 방문'}"
        )
        add_table(
            doc,
            ["역할", "이름", "이메일", "연락처(카톡ID)", "대학", "링크드인", "비고"],
            [
                [
                    m["role"],
                    m["name"],
                    m["email"],
                    m["contact"],
                    m.get("university") or "",
                    m.get("linkedin") or "-",
                    "",
                ]
                for m in r["members"]
            ],
        )

    # ── 5. 일괄 발송용 목록
    doc.add_paragraph("5. 일괄 발송용 목록", style="Heading 1")
    doc.add_paragraph("중복은 이미 제외된 명단이라, 그대로 붙여 넣어도 같은 사람에게 두 번 가지 않습니다.")

    emails, seen = [], set()
    for p in people:
        k = (p.get("email") or "").strip().lower()
        if k and k not in seen:
            seen.add(k)
            emails.append(p["email"])
    doc.add_paragraph("이메일 (쉼표 구분)", style="Heading 2")
    doc.add_paragraph(", ".join(emails))
    doc.add_paragraph(f"고유 이메일 {len(emails)}개")

    contacts, seen = [], set()
    for p in people:
        k = (p.get("contact") or "").strip().lower()
        if k and k not in seen:
            seen.add(k)
            contacts.append(p["contact"])
    doc.add_paragraph("연락처 / 카카오톡 ID (쉼표 구분)", style="Heading 2")
    doc.add_paragraph(", ".join(contacts))
    doc.add_paragraph(f"고유 연락처 {len(contacts)}개")

    doc.save(str(OUT))
    print(
        f"{OUT.name}: 신청 {len(regs)}건 / 명단 {len(people)}명 / 중복 제외 {len(dropped)}명"
    )


if __name__ == "__main__":
    main()
